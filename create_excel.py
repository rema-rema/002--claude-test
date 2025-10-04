#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel file creation script for test.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# Create workbook
wb = Workbook()
ws = wb.active

# A1にあああを入力
ws['A1'] = 'あああ'

# B1～F2の範囲に表を作成（ヘッダー + データ行）
headers = ['項目1', '項目2', '項目3', '項目4', '項目5']
data = [100, 200, 300, 400, 500]

# ヘッダー行（B1～F1）
for idx, header in enumerate(headers, start=2):  # B列は2番目
    cell = ws.cell(row=1, column=idx)
    cell.value = header
    # ヘッダーのスタイル（太字、背景色、中央揃え）
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')

# データ行（B2～F2）
for idx, value in enumerate(data, start=2):  # B列は2番目
    cell = ws.cell(row=2, column=idx)
    cell.value = value
    # データ行のスタイル（中央揃え、背景色）
    cell.fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center')

# 罫線を設定（B1～F2の範囲）
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for row in range(1, 3):  # 1行目～2行目
    for col in range(2, 7):  # B列(2)～F列(6)
        ws.cell(row=row, column=col).border = thin_border

# 列幅を調整
ws.column_dimensions['A'].width = 12
for col in ['B', 'C', 'D', 'E', 'F']:
    ws.column_dimensions[col].width = 15

# 保存
output_path = '/home/rema/project/002--claude-test/test.xlsx'
wb.save(output_path)

print(f"Excel file created: {output_path}")
